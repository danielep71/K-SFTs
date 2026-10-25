"""Headless regression of the K-SFTs engine in LibreOffice (VBA compatibility mode).

Not a substitute for Excel: LibreOffice runs the exported VBA with
`Option VBASupport 1`, which is close to, but not, the VBE. It is the check
that can run without Excel (CI, Linux): it loads the template, inserts the
modules of src/ into the document's Basic library, runs CORE_Engine.Calculate
and every SFT_* unit test, and compares Results with
tests/derivations/prototype/demo-portfolio.expected.json.

Harness-only shims, applied to an in-memory copy of the source, never to the
files: `vbObjectError + n` constants are replaced by their value (LibreOffice
does not accept that constant expression), `Option Private Module` is removed,
`Application.CalculationState <> xlDone` becomes False and `Debug.Print` is
dropped (neither exists in LibreOffice).

Usage: python tools/lo_regression.py [--keep OUT.xlsx]
Needs LibreOffice (soffice) and its Python UNO bridge.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped,import-not-found,unused-ignore]
import uno  # type: ignore[import-not-found,unused-ignore]
from com.sun.star.beans import PropertyValue  # type: ignore[import-not-found,unused-ignore]

ROOT = Path(__file__).resolve().parents[1]
MODULES = ["src/core/CORE_Config.bas", "src/core/CORE_Util.bas", "src/modules/SFT_Formulas.bas",
           "src/core/CORE_Engine.bas", "src/workbook/M_Main.bas"]
RESULT_COLUMNS = {7: "holding_bd", 10: "sum_e", 11: "sum_c", 12: "net", 13: "gross", 14: "n", 15: "systematic",
                  16: "idiosyncratic", 17: "fx_term", 18: "e_star", 20: "rw_key", 21: "rw", 22: "rwa",
                  23: "own_funds", 24: "status"}


def prop(name: str, value: Any) -> Any:
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


def shim(source: str) -> str:
    lines = source.replace("\r\n", "\n").split("\n")
    source = "\n".join(line for line in lines if not line.startswith("Attribute VB_"))
    source = re.sub(r"vbObjectError \+ (\d+)", lambda m: str(-2147221504 + int(m.group(1))), source)
    source = source.replace("Option Private Module", "")
    source = source.replace("Application.CalculationState <> xlDone", "False")
    source = re.sub(r"^\s*Debug\.Print .*$", "", source, flags=re.M)
    return "Option VBASupport 1\n" + source


def formula_tests() -> tuple[str, int]:
    """Basic module running the FormulaTest rows of tools/build_template.py."""
    text = (ROOT / "tools/build_template.py").read_text()
    block = text[text.index("TESTS = ["):text.index("]\nr = 6\nfor k, (formula") + 1]
    tests = eval(block.split("=", 1)[1])  # the list literal written in build_template.py
    lines = ["Option Explicit", "Public Function RunFormulaTests() As String", "    Dim v As Variant",
             "    Dim passed As Long", "    Dim failed As String"]
    for k, (formula, expected, _) in enumerate(tests, 1):
        call = re.sub(r"\{([^}]*)\}", lambda m: "Array(" + m.group(1) + ")", formula)
        call = call.replace("TRUE", "True").replace("FALSE", "False")
        lines.append(f"    v = {call}")
        if expected == "#NUM!":
            lines.append(f'    If IsError(v) Then\n        If CLng(v) = 2036 Then passed = passed + 1 Else '
                         f'failed = failed & " F-{k:02d}"\n    Else\n        failed = failed & " F-{k:02d}"\n    End If')
        else:
            lines.append(f'    If IsError(v) Then\n        failed = failed & " F-{k:02d}"\n    ElseIf Abs(v - '
                         f'({expected!r})) <= 1E-9 Then\n        passed = passed + 1\n    Else\n        '
                         f'failed = failed & " F-{k:02d}"\n    End If')
    lines += [f'    RunFormulaTests = passed & " of {len(tests)} PASS" & failed', "End Function",
              "Public Function RunEngine() As String",
              "    RunEngine = M_Main.RunSFT_Silent()", "End Function"]
    return "\n".join(lines) + "\n", len(tests)


def run(out_path: Path) -> list[str]:
    port = 2002
    profile = Path(tempfile.mkdtemp(prefix="ksft_lo_"))
    office = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore", "--nologo",
                               f"--accept=socket,host=localhost,port={port};urp;",
                               f"-env:UserInstallation={profile.as_uri()}"],
                              stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    local = uno.getComponentContext()
    resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
    for _ in range(120):
        try:
            ctx = resolver.resolve(f"uno:socket,host=localhost,port={port};urp;StarOffice.ComponentContext")
            break
        except Exception:
            time.sleep(0.5)
    else:
        raise RuntimeError("LibreOffice did not start")
    desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
    template = uno.systemPathToFileUrl(str(ROOT / "src/workbook/SFT_Template.xlsx"))
    doc = desktop.loadComponentFromURL(template, "_blank", 0, (prop("Hidden", True), prop("MacroExecutionMode", 4)))
    library = doc.BasicLibraries.getByName("Standard") if doc.BasicLibraries.hasByName("Standard") \
        else doc.BasicLibraries.createLibrary("Standard")
    for module in MODULES:
        source = (ROOT / module).read_text(encoding="cp1252")
        found = re.search(r'Attribute VB_Name = "([^"]+)"', source)
        if found is None:
            raise ValueError(f"{module}: no Attribute VB_Name line")
        name = found.group(1)
        library.insertByName(name, shim(source))
    test_source, _ = formula_tests()
    library.insertByName("TEST_LO", "Option VBASupport 1\n" + test_source)
    provider = doc.getScriptProvider()
    outcomes = []
    for proc in ("RunEngine", "RunFormulaTests"):
        script = provider.getScript(f"vnd.sun.star.script:Standard.TEST_LO.{proc}?language=Basic&location=document")
        outcomes.append(str(script.invoke((), (), ())[0]))
    doc.calculateAll()
    doc.storeToURL(uno.systemPathToFileUrl(str(out_path)), (prop("FilterName", "Calc Office Open XML"),))
    doc.close(True)
    try:
        desktop.terminate()
    except Exception:
        pass
    office.wait(timeout=60)
    return outcomes


def compare(out_path: Path) -> tuple[int, list[str]]:
    expected = json.loads((ROOT / "tests/derivations/prototype/demo-portfolio.expected.json").read_text())["results"]
    sheet = openpyxl.load_workbook(out_path, data_only=True)["Results"]
    checks, problems, seen = 0, [], set()
    for row in range(7, 7 + 2000):
        nsid = sheet.cell(row, 3).value
        if not nsid:
            break
        if nsid not in expected:
            continue
        if nsid in seen:
            problems.append(f"{nsid}: more than one Results row")
            continue
        seen.add(nsid)
        for column, key in RESULT_COLUMNS.items():
            if key not in expected[nsid]:
                continue
            want, got = expected[nsid][key], sheet.cell(row, 2 + column).value
            checks += 1
            same = abs((got or 0) - want) <= 1e-6 * max(1.0, abs(want)) if isinstance(want, (int, float)) \
                else got == want
            if not same:
                problems.append(f"{nsid} {key}: calculated {got!r}, expected {want!r}")
    problems += [f"{nsid}: no Results row" for nsid in sorted(set(expected) - seen)]
    if checks == 0:
        problems.append("no Results figure compared")
    return checks, problems


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--keep", type=Path, help="keep the calculated workbook at this path")
    args = parser.parse_args()
    out_path = (args.keep or Path(tempfile.mkdtemp()) / "k-sft-lo.xlsx").resolve()
    engine, formulas = run(out_path)
    checks, problems = compare(out_path)
    print("engine:  ", engine)
    print("formulas:", formulas)
    print(f"results:  {checks} figures compared, {len(problems)} differ")
    for p in problems:
        print("   ", p)
    ok = engine.startswith("OK") and " PASS" in formulas and formulas.split(" of ")[0] == \
        formulas.split(" of ")[1].split(" ")[0] and not problems
    print("REGRESSION", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
